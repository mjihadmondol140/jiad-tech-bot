import os
import json
import asyncio
from datetime import datetime, time
from zoneinfo import ZoneInfo
import io
import html

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, ReplyKeyboardMarkup, KeyboardButton, CopyTextButton
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)
import openpyxl
from playwright.sync_api import sync_playwright

# ================= CONFIG =================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

with open(os.path.join(BASE_DIR, "config.txt"), "r", encoding="utf-8") as f:
    BOT_TOKEN = f.read().strip()

DEFAULT_CHANNEL = "@OMORTECH"
DEFAULT_ADMIN_ID = 7815818738

DEFAULT_TARGET_CHAT_IDS = [
    6945991387,
    7815818738,
]

DB_FILE = os.path.join(BASE_DIR, "bot_data.json")
TZ = ZoneInfo("Asia/Dhaka")


# ================= DATABASE & SETTINGS =================

def load_db():
    default_data = {
        "settings": {
            "channels": [DEFAULT_CHANNEL],
            "admin_ids": [DEFAULT_ADMIN_ID],
            "target_chat_ids": DEFAULT_TARGET_CHAT_IDS,
            "maintenance": False,
            "approval_mode": True,
            "min_ids_limit": 1,
            "duplicate_mode": "daily",
            "price_only_uid": 1.5,
            "price_page_uid": 2.5,
            "payment_msg_template": "🎉 <b>আপনার পেমেন্ট সফলভাবে পরিশোধ করা হয়েছে!</b>\n\n💵 <b>পরিশোধিত অর্থ:</b> {amount} BDT",
            "work_schedule_enabled": False,
            "start_time": "10:00",
            "end_time": "22:00",
            "send_live_check_file": True
        },
        "payment_methods": ["bKash", "Nagad", "Rocket", "Binance"],
        "user_payment_details": {},
        "user_languages": {},
        "banned_users": [],
        "pending_submissions": {},
        "days": {},
        "payments": {}
    }

    if not os.path.exists(DB_FILE):
        return default_data

    try:
        with open(DB_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)

        data.setdefault("settings", {})
        
        # Backward compatibility migration for channels & admins
        if "channel" in data["settings"] and "channels" not in data["settings"]:
            ch = data["settings"].pop("channel")
            data["settings"]["channels"] = [ch] if ch else [DEFAULT_CHANNEL]
        if "admin_id" in data["settings"] and "admin_ids" not in data["settings"]:
            aid = data["settings"].pop("admin_id")
            data["settings"]["admin_ids"] = [aid] if aid else [DEFAULT_ADMIN_ID]

        for k, v in default_data["settings"].items():
            data["settings"].setdefault(k, v)
        
        data.setdefault("payment_methods", ["bKash", "Nagad", "Rocket", "Binance"])
        if "Binance" not in data["payment_methods"]:
            data["payment_methods"].append("Binance")

        data.setdefault("user_payment_details", {})
        data.setdefault("user_languages", {})
        data.setdefault("banned_users", [])
        data.setdefault("pending_submissions", {})
        data.setdefault("days", {})
        data.setdefault("payments", {})

        return data
    except (json.JSONDecodeError, OSError):
        return default_data


DATA = load_db()


def _save_db_file():
    temp_file = DB_FILE + ".tmp"
    with open(temp_file, "w", encoding="utf-8") as f:
        json.dump(DATA, f, ensure_ascii=False, indent=2)
    os.replace(temp_file, DB_FILE)


def save_db():
    asyncio.create_task(asyncio.to_thread(_save_db_file))


def today():
    return datetime.now(TZ).strftime("%Y-%m-%d")


def is_bot_admin(user_id):
    admin_ids = DATA["settings"].get("admin_ids", [DEFAULT_ADMIN_ID])
    return user_id in admin_ids


def is_work_time():
    s = DATA["settings"]
    if not s.get("work_schedule_enabled", False):
        return True
    
    now_time = datetime.now(TZ).time()
    try:
        sh, sm = map(int, s.get("start_time", "10:00").split(":"))
        eh, em = map(int, s.get("end_time", "22:00").split(":"))
        start_t = time(sh, sm)
        end_t = time(eh, em)
        
        if start_t <= end_t:
            return start_t <= now_time <= end_t
        else:
            return now_time >= start_t or now_time <= end_t
    except Exception:
        return True


def get_day(date):
    if date not in DATA["days"]:
        DATA["days"][date] = {
            "ids": [],
            "members": {},
        }
    day = DATA["days"][date]
    day.setdefault("ids", [])
    day.setdefault("members", {})
    return day


def get_member(day, user):
    uid = str(user.id)
    members = day["members"]

    if uid not in members:
        members[uid] = {
            "name": user.full_name,
            "received": 0,
            "new": 0,
            "duplicate": 0,
            "rejected": 0,
            "only_uid_count": 0,
            "page_uid_count": 0,
            "all_ids": [],
            "new_ids": [],
            "duplicate_ids": [],
            "rejected_ids": []
        }

    member = members[uid]
    member["name"] = user.full_name
    member.setdefault("rejected", 0)
    member.setdefault("rejected_ids", [])
    member.setdefault("only_uid_count", 0)
    member.setdefault("page_uid_count", 0)
    return member


def calculate_member_earned(member):
    p_only = DATA["settings"].get("price_only_uid", 0.0)
    p_page = DATA["settings"].get("price_page_uid", 0.0)
    only_c = member.get("only_uid_count", 0)
    page_c = member.get("page_uid_count", 0)
    
    total_earned = (only_c * p_only) + (page_c * p_page)
    return max(0.0, total_earned)


def get_user_payment_status(uid, date):
    return DATA["payments"].get(f"{uid}_{date}", "Unpaid")


# ================= LIVE FB UID CHECK FUNCTION =================

def check_facebook_uid(uids_input):
    if isinstance(uids_input, list):
        uids_text = "\n".join(uids_input)
    else:
        uids_text = str(uids_input).strip()

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(
                headless=True,
                args=[
                    "--no-sandbox",
                    "--disable-setuid-sandbox",
                    "--disable-dev-shm-usage",
                    "--disable-gpu"
                ]
            )
            context = browser.new_context(
                user_agent=(
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
                    " (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
                )
            )
            page = context.new_page()

            page.goto("https://uid.sure-sms.com/", timeout=30000)
            page.wait_for_selector("textarea", timeout=10000)

            page.fill("textarea", uids_text)

            try:
                page.click("button:has-text('Check')", timeout=3000)
            except:
                try:
                    page.click("button:has-text('Submit')", timeout=3000)
                except:
                    page.click("button", timeout=3000)

            page.wait_for_timeout(4000)

            textareas = page.locator("textarea").all()
            live_uids = []
            die_uids = []

            if len(textareas) >= 3:
                live_text = textareas[1].input_value()
                die_text = textareas[2].input_value()

                live_uids = [u.strip() for u in live_text.split("\n") if u.strip()]
                die_uids = [u.strip() for u in die_text.split("\n") if u.strip()]
            elif len(textareas) == 2:
                live_text = textareas[1].input_value()
                live_uids = [u.strip() for u in live_text.split("\n") if u.strip()]

            browser.close()

            input_set = set(u.strip() for u in uids_text.split("\n") if u.strip())
            valid_live = [u for u in live_uids if u in input_set]
            valid_die = [u for u in die_uids if u in input_set]
            
            accounted = set(valid_live + valid_die)
            for u in input_set:
                if u not in accounted:
                    if u not in valid_die:
                        valid_die.append(u)

            return {"live": valid_live, "die": valid_die}

    except Exception as e:
        print(f"Fast Check Error: {e}")
        input_list = [u.strip() for u in uids_text.split("\n") if u.strip()]
        return {"live": input_list, "die": []}
# ================= CHANNEL JOIN & BAN CHECK =================

async def is_joined(user_id, context):
    if is_bot_admin(user_id):
        return True

    channels = DATA["settings"].get("channels", [DEFAULT_CHANNEL])
    if not channels:
        return True
        
    for channel in channels:
        try:
            member = await context.bot.get_chat_member(
                chat_id=channel,
                user_id=user_id,
            )
            if member.status not in ("member", "administrator", "creator"):
                return False
        except Exception:
            pass
    return True


async def require_join(update, context):
    channels = DATA["settings"].get("channels", [DEFAULT_CHANNEL])
    keyboard = []
    for idx, ch in enumerate(channels, 1):
        url = f"https://t.me/{ch.replace('@', '')}" if ch.startswith("@") else ch
        keyboard.append([InlineKeyboardButton(f"📢 Join Channel {idx}", url=url)])

    await update.effective_message.reply_text(
        "🚫 সার্ভিসটি ব্যবহার করতে হলে প্রথমে আমাদের চ্যানেলগুলোতে জয়েন করতে হবে。\n\n"
        "👉 চ্যানেলগুলোতে জয়েন করার পর /start দিন。",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )


# ================= START, LANGUAGE & STATUS =================

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    user = update.effective_user

    if user.id in DATA["banned_users"]:
        await update.effective_message.reply_text("🚫 আপনাকে সিস্টেম থেকে ব্লক করা হয়েছে।")
        return

    if DATA["settings"]["maintenance"] and not is_bot_admin(user.id):
        await update.effective_message.reply_text("🛠️ বট বর্তমানে মেনটেইনেন্স মোডে আছে। কিছুক্ষণ পর আবার চেষ্টা করুন।")
        return

    if not await is_joined(user.id, context):
        await require_join(update, context)
        return

    date = today()
    day = get_day(date)
    get_member(day, user)
    save_db()

    uid_str = str(user.id)
    pay_details = DATA["user_payment_details"].get(uid_str, {})
    if isinstance(pay_details, dict) and pay_details:
        pay_lines = [f"• {m}: <code>{n}</code>" for m, n in pay_details.items() if n]
        pay_info = "\n".join(pay_lines) if pay_lines else "নম্বর সেট করা হয়নি"
    else:
        pay_info = "নম্বর সেট করা হয়নি"

    lang = DATA["user_languages"].get(uid_str, "bn")

    if lang == "en":
        msg = (
            "<b>🤖 JIHAD TECH SUBMISSION BOT</b>\n\n"
            f"👤 <b>Name:</b> {html.escape(user.full_name)}\n"
            f"🆔 <b>User ID:</b> <code>{user.id}</code>\n"
            f"💳 <b>Payment Number:</b>\n{pay_info}\n\n"
            "🟢 Bot is active. Send your ID list or <code>.xlsx</code> file.\n\n"
            "📊 কাজের হিসাব দেখতে: /status\n"
            "📋 জমা দেওয়া আইডি: /myids\n"
            "💳 পেমেন্ট নম্বর সেট করতে: /setpayment\n"
            "🌐 ভাষা পরিবর্তন করতে: /language\n"
            "🎫 সাহায্য পেতে: /support"
        )
    else:
        msg = (
            "<b>🤖 JIHAD TECH RECEVER BOT</b>\n\n"
            f"👤 <b>Name:</b> {html.escape(user.full_name)}\n"
            f"🆔 <b>User ID:</b> <code>{user.id}</code>\n"
            f"💳 <b>Payment Number:</b>\n{pay_info}\n\n"
            "🟢 বট সক্রিয় আছে। আপনার প্রয়োজনীয় সংখ্যার ID list অথবা<code>.xlsx</code> ফাইল পাঠান。\n\n"
            "📊 কাজের হিসাব দেখতে: /status\n"
            "📋 জমা দেওয়া আইডি: /myids\n"
            "💳 পেমেন্ট নম্বর সেট করতে: /setpayment\n"
            "🌐 ভাষা পরিবর্তন করতে: /language\n"
            "🎫 সাহায্য পেতে: /support"
        )

    reply_markup = None
    if is_bot_admin(user.id):
        reply_markup = ReplyKeyboardMarkup(
            [[KeyboardButton("👑 অ্যাডমিন প্যানেল")]],
            resize_keyboard=True
        )

    await update.effective_message.reply_text(msg, parse_mode="HTML", reply_markup=reply_markup)


async def my_ids_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if user.id in DATA["banned_users"]:
        return
    if not await is_joined(user.id, context):
        await require_join(update, context)
        return

    date = today()
    uid = str(user.id)
    day_data = DATA["days"].get(date)

    if not day_data or uid not in day_data["members"] or not day_data["members"][uid]["all_ids"]:
        await update.effective_message.reply_text(
            f"📅 <b>Date:</b> <code>{date}</code>\n\n📭 আজকের দিনে আপনি এখনো কোনো আইডি জমা দেননি।",
            parse_mode="HTML"
        )
        return

    mem = day_data["members"][uid]
    user_ids = mem.get("all_ids", [])

    def gen_my_excel():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "My IDs"
        ws.append(["UID", "Date"])
        for u in user_ids:
            ws.append([u, date])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return bio

    try:
        excel_bio = await asyncio.to_thread(gen_my_excel)
        await update.effective_message.reply_document(
            document=excel_bio,
            filename=f"my_submitted_ids_{date}.xlsx",
            caption=f"📋 আপনার আজকের জমা দেওয়া মোট {len(user_ids)} টি আইডির এক্সেল ফাইল।"
        )
    except Exception as e:
        print("Failed to send my_ids excel:", e)


async def language_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    keyboard = [
        [
            InlineKeyboardButton("🇧🇩 বাংলা", callback_data="setlang_bn"),
            InlineKeyboardButton("🇬🇧 English", callback_data="setlang_en")
        ]
    ]
    await update.effective_message.reply_text(
        "🌐 <b>Select Your Preferred Language / আপনার পছন্দের ভাষা নির্বাচন করুন:</b>",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="HTML"
    )


async def set_payment(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    methods = DATA.get("payment_methods", ["bKash", "Nagad", "Rocket", "Binance"])
    
    keyboard = []
    for m in methods:
        keyboard.append([InlineKeyboardButton(f"💳 {m}", callback_data=f"setpaym_{m}")])

    await update.effective_message.reply_text(
        "💳 <b>পেমেন্ট মেথড সিলেক্ট করুন:</b>",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="HTML"
    )


async def support_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["waiting_for"] = "user_support_msg"
    await update.effective_message.reply_text(
        "📩 <b>SUPPORT & HELP DESK</b>\n\n"
        "আপনার কোনো প্রশ্ন, সমস্যা বা পেমেন্ট সংক্রান্ত বিষয় থাকলে নিচে বিস্তারিত লিখে মেসেজ পাঠান। অ্যাডমিন আপনার বার্তাটি পেয়ে সরাসরি উত্তর দেবেন।",
        parse_mode="HTML"
    )


async def status(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user

    if not await is_joined(user.id, context):
        await require_join(update, context)
        return

    uid = str(user.id)
    dates = sorted(DATA["days"].keys(), reverse=True)
    today_date = today()

    lines = [
        "<b>📊 YOUR SUBMISSION STATUS</b>",
        f"👤 <b>Name:</b> {html.escape(user.full_name)}",
        f"🆔 <b>User ID:</b> <code>{user.id}</code>\n",
    ]

    total_received = total_new = total_dup = total_rej = 0
    total_only_uid = total_page_uid = 0
    total_earned = 0.0
    found = False
    shown_days = 0

    for date in dates:
        member = DATA["days"][date]["members"].get(uid)
        if not member:
            continue

        found = True
        rej_count = member.get("rejected", 0)
        total_received += member["received"]
        total_new += member["new"]
        total_dup += member["duplicate"]
        total_rej += rej_count
        
        o_cnt = member.get("only_uid_count", 0)
        p_cnt = member.get("page_uid_count", 0)
        total_only_uid += o_cnt
        total_page_uid += p_cnt

        earned = calculate_member_earned(member)
        total_earned += earned

        if shown_days < 7:
            pay_st = get_user_payment_status(uid, date)
            tag = " 🌟 (Today)" if date == today_date else ""
            lines.extend([
                f"📅 <b>Date:</b> <code>{date}</code>{tag}",
                f"📥 Total Received: <code>{member['received']}</code>",
                f"🟢 New IDs: <code>{member['new']}</code> (Only: {o_cnt} | Page: {p_cnt})",
                f"♻️ Duplicate IDs: <code>{member['duplicate']}</code>",
                f"🚫 Rejected IDs: <code>{rej_count}</code>",
                f"💰 Earned Amount: <code>{earned:.2f} BDT</code> ({pay_st})",
                "────────────────────\n"
            ])
            shown_days += 1

    if not found:
        lines.append("📭 আপনি এখনো কোনো আইডি জমা দেননি।")
        await update.effective_message.reply_text("\n".join(lines), parse_mode="HTML")
        return

    lines.extend([
        "━━━━━━━━━━━━━━━━━━━━",
        "📦 <b>ALL-TIME SUMMARY</b>",
        f"📥 Total Submitted: <code>{total_received}</code>",
        f"🟢 Total New: <code>{total_new}</code> (Only: {total_only_uid} | Page: {total_page_uid})",
        f"♻️ Total Duplicate: <code>{total_dup}</code>",
        f"🚫 Total Rejected: <code>{total_rej}</code>",
        f"💰 Total Earned: <code>{total_earned:.2f} BDT</code>",
    ])

    await update.effective_message.reply_text("\n".join(lines), parse_mode="HTML")


# ================= ID & FILE PROCESSING =================

async def process_valid_submission(user, ids, category, context, update_msg, file_info=None):
    date = today()
    day = get_day(date)
    member = get_member(day, user)

    mode = DATA["settings"]["duplicate_mode"]
    if mode == "all_time":
        all_existing_ids = set()
        for d_data in DATA["days"].values():
            all_existing_ids.update(d_data["ids"])
        seen_today = all_existing_ids
    else:
        seen_today = set(day["ids"])

    seen_message = set()
    new_ids = []
    duplicate_ids = []

    for item in ids:
        if item in seen_today or item in seen_message:
            duplicate_ids.append(item)
        else:
            new_ids.append(item)
        seen_message.add(item)

    day["ids"].extend(ids)
    member["received"] += len(ids)
    member["new"] += len(new_ids)
    member["duplicate"] += len(duplicate_ids)
    
    if category == "ONLY UID":
        member["only_uid_count"] += len(new_ids)
    else:
        member["page_uid_count"] += len(new_ids)

    member["all_ids"].extend(ids)
    member["new_ids"].extend(new_ids)
    member["duplicate_ids"].extend(duplicate_ids)
    save_db()

    try:
        await context.bot.send_message(
            chat_id=user.id,
            text=(
                f"🟢 <b>আপনার সাবমিশন অনুমোদিত (Approved) হয়েছে!</b>\n\n"
                f"🏷️ ক্যাটাগরি: {category}\n"
                f"📦 মোট আইডি: {len(ids)}\n"
                f"🟢 নতুন আইডি: {len(new_ids)} টি\n"
                f"♻️ ডুপ্লিকেট আইডি: {len(duplicate_ids)} টি"
            ),
            parse_mode="HTML"
        )
    except Exception:
        pass

    if new_ids:
        chunk_size = 150
        for i in range(0, len(new_ids), chunk_size):
            chunk = new_ids[i:i + chunk_size]
            chunk_text = "\n".join(chunk)
            
            report = (
                "<b>📋 NEW APPROVED SUBMISSION</b>\n\n"
                f"👤 Name: <a href='tg://user?id={user.id}'>{html.escape(user.full_name)}</a>\n"
                f"🆔 <b>User ID:</b> <code>{user.id}</code>\n"
                f"🏷️ <b>Category:</b> {category}\n"
                f"📅 <b>Date:</b> {date}\n\n"
                f"🟢 <b>New IDs ({len(chunk)}):</b>\n"
                f"<pre>{chunk_text}</pre>"
            )

            keyboard = [
                [
                    InlineKeyboardButton(
                        f"📋 এই ব্যাচের {len(chunk)} টি আইডি কপি করুন", 
                        copy_text=CopyTextButton(text=chunk_text)
                    )
                ]
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)

            for chat_id in DATA["settings"]["target_chat_ids"]:
                try:
                    await context.bot.send_message(
                        chat_id=chat_id,
                        text=report,
                        reply_markup=reply_markup,
                        parse_mode="HTML"
                    )
                    await asyncio.sleep(0.1)
                except Exception as e:
                    print(f"Message send failed for {chat_id}: {e}")

        if file_info and file_info.get("document_id"):
            for chat_id in DATA["settings"]["target_chat_ids"]:
                try:
                    await context.bot.send_document(
                        chat_id=chat_id,
                        document=file_info["document_id"],
                        caption=f"📁 File from {user.full_name} ({date}) [{category}]"
                    )
                    await asyncio.sleep(0.1)
                except Exception:
                    pass


async def handle_incoming_data(user, ids, category, context, update_msg, file_info=None):
    check_msg = await update_msg.reply_text("🔍 UID গুলো লাইভ চেক করা হচ্ছে, অনুগ্রহ করে অপেক্ষা করুন...")
    
    check_result = await asyncio.to_thread(check_facebook_uid, ids)
    live_uids = check_result.get("live", [])
    die_uids = check_result.get("die", [])
    
    valo_count = len(live_uids)
    nosto_count = len(die_uids)

    min_limit = DATA["settings"]["min_ids_limit"]
    if valo_count < min_limit:
        try:
            await check_msg.edit_text(
                f"⚠️ সর্বনিম্ন <code>{min_limit}</code>টি <b>Live</b> আইডি একসঙ্গে পাঠাতে হবে। "
                f"আপনি মোট দিয়েছেন <code>{len(ids)}</code>টি, যার মধ্যে <b>Live</b> হয়েছে <code>{valo_count}</code>টি এবং <b>Dead</b> হয়েছে <code>{nosto_count}</code>টি。",
                parse_mode="HTML"
            )
        except Exception:
            pass
        return

    # Send Live Check Excel File FIRST with caption together, status column as Live / Dead
    if DATA["settings"].get("send_live_check_file", True):
        def gen_check_excel():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Live Check Results"
            ws.append(["UID", "Status"])
            for u in live_uids:
                ws.append([u, "Live"])
            for u in die_uids:
                ws.append([u, "Dead"])
            bio = io.BytesIO()
            wb.save(bio)
            bio.seek(0)
            return bio

        try:
            excel_bio = await asyncio.to_thread(gen_check_excel)
            if DATA["settings"]["approval_mode"]:
                caption_text = (
                    f"✅ <b>আপনার পাঠানো আইডি সফলভাবে লাইভ চেক সম্পন্ন হয়েছে!</b>\n\n"
                    f"🟢 Live: {valo_count} | 🔴 Dead: {nosto_count}\n"
                    f"⏳ আপনার সাবমিশনটি বর্তমানে অ্যাডমিনের অনুমোদনের জন্য পেন্ডিং রয়েছে।\n"
                    f"📁 লাইভ চেক করা আইডিসমূহের এক্সেল ফাইল।"
                )
            else:
                caption_text = (
                    f"✅ <b>আপনার পাঠানো আইডি সফলভাবে চেক ও জমা হয়েছে!</b>\n\n"
                    f"🟢 Live: {valo_count} | 🔴 Dead: {nosto_count}\n"
                    f"📁 লাইভ চেক করা আইডিসমূহের এক্সেল ফাইল।"
                )
            await update_msg.reply_document(
                document=excel_bio,
                filename=f"live_check_{today()}.xlsx",
                caption=caption_text,
                parse_mode="HTML"
            )
        except Exception as e:
            print("Failed to send live check excel:", e)

    valid_ids_to_process = live_uids

    date = today()
    day = get_day(date)
    mode = DATA["settings"]["duplicate_mode"]
    if mode == "all_time":
        all_existing_ids = set()
        for d_data in DATA["days"].values():
            all_existing_ids.update(d_data["ids"])
        seen_today = all_existing_ids
    else:
        seen_today = set(day["ids"])

    seen_msg = set()
    temp_new = 0
    temp_dup = 0
    for item in valid_ids_to_process:
        if item in seen_today or item in seen_msg:
            temp_dup += 1
        else:
            temp_new += 1
        seen_msg.add(item)

    if DATA["settings"]["approval_mode"]:
        sub_id = f"sub_{user.id}_{int(datetime.now().timestamp())}"
        DATA["pending_submissions"][sub_id] = {
            "user_id": user.id,
            "full_name": user.full_name,
            "ids": valid_ids_to_process,
            "category": category,
            "file_info": file_info
        }
        save_db()

        # Send notification to ALL configured admins
        admin_ids = DATA["settings"].get("admin_ids", [DEFAULT_ADMIN_ID])
        admin_preview_name = f"<a href='tg://user?id={user.id}'>{html.escape(user.full_name)}</a>"
        
        live_len = len(valid_ids_to_process)
        preview_text = (
            "<b>🔔 NEW SUBMISSION PENDING</b>\n\n"
            f"👤 Name: {admin_preview_name}\n"
            f"🆔 <b>User ID:</b> <code>{user.id}</code>\n"
            f"🏷️ <b>Category:</b> {category}\n"
            f"🟢 <b>Live:</b> <code>{live_len}</code> | 🔴 <b>Dead:</b> <code>{nosto_count}</code>\n"
            f"📥 <b>New:</b> <code>{temp_new}</code> | ♻️ <b>Dupes:</b> <code>{temp_dup}</code>\n\n"
            "নিচের বাটন থেকে এক-ক্লিকে প্রসেস করুন:"
        )

        keyboard = [
            [InlineKeyboardButton("🟢 Accept New Only (Skip Dupes)", callback_data=f"acc_{sub_id}")],
            [InlineKeyboardButton("🚫 Reject Entire Batch", callback_data=f"rej_{sub_id}")]
        ]
        
        for adm_id in admin_ids:
            try:
                await context.bot.send_message(
                    chat_id=adm_id,
                    text=preview_text,
                    reply_markup=InlineKeyboardMarkup(keyboard),
                    parse_mode="HTML"
                )
            except Exception as e:
                print(f"Failed to send notification to admin {adm_id}:", e)

        try:
            await check_msg.delete()
        except Exception:
            pass
    else:
        try:
            await check_msg.delete()
        except Exception:
            pass
        await process_valid_submission(user, valid_ids_to_process, category, context, update_msg, file_info)


async def handle_admin_inputs(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    waiting = context.user_data.get("waiting_for")
    if not waiting:
        return False

    text = update.effective_message.text.strip()

    if waiting == "user_enter_pay_num":
        context.user_data["waiting_for"] = None
        selected_method = context.user_data.get("selected_pay_method", "bKash")
        uid_str = str(user.id)
        
        if uid_str not in DATA["user_payment_details"] or not isinstance(DATA["user_payment_details"][uid_str], dict):
            DATA["user_payment_details"][uid_str] = {}
            
        DATA["user_payment_details"][uid_str][selected_method] = text
        save_db()
        await update.effective_message.reply_text(f"✅ আপনার <b>{selected_method}</b> নম্বর সফলভাবে সেভ করা হয়েছে: <code>{text}</code>", parse_mode="HTML")
        return True

    if waiting == "user_support_msg":
        context.user_data["waiting_for"] = None
        admin_ids = DATA["settings"].get("admin_ids", [DEFAULT_ADMIN_ID])
        
        supp_text = (
            "<b>📩 NEW SUPPORT TICKET RECEIVED</b>\n\n"
            f"👤 Name: {html.escape(user.full_name)}\n"
            f"🆔 <b>User ID:</b> <code>{user.id}</code>\n\n"
            f"<b>💬 Message:</b>\n{text}"
        )
        for adm_id in admin_ids:
            try:
                await context.bot.send_message(chat_id=adm_id, text=supp_text, parse_mode="HTML")
            except Exception:
                pass
        await update.effective_message.reply_text("✅ আপনার মেসেজটি অ্যাডমিনের কাছে সফলভাবে পাঠানো হয়েছে। অতি শীঘ্রই উত্তর পেয়ে যাবেন।")
        return True

    if not is_bot_admin(user.id):
        return False

    if waiting == "msg_user_id":
        context.user_data["waiting_for"] = "msg_user_text"
        context.user_data["target_msg_uid"] = text
        await update.effective_message.reply_text(
            f"✉️ ইউজার আইডি <code>{text}</code> সিলেক্ট হয়েছে। এবার ইউজারকে যে মেসেজটি পাঠাতে চান তা লিখে পাঠান:",
            parse_mode="HTML"
        )
        return True

    elif waiting == "msg_user_text":
        context.user_data["waiting_for"] = None
        target_uid = context.user_data.get("target_msg_uid")
        try:
            await context.bot.send_message(
                chat_id=int(target_uid),
                text=f"📬 <b>ADMIN MESSAGE</b>\n\n{text}",
                parse_mode="HTML"
            )
            await update.effective_message.reply_text(f"✅ সফলভাবে ইউজার <code>{target_uid}</code> এর কাছে মেসেজ পাঠানো হয়েছে।", parse_mode="HTML")
        except Exception as e:
            await update.effective_message.reply_text(f"🚫 মেসেজ পাঠানো যায়নি: {e}")
        return True

    elif waiting == "add_channel":
        context.user_data["waiting_for"] = None
        channels = DATA["settings"].setdefault("channels", [])
        if text not in channels:
            channels.append(text)
            save_db()
            await update.effective_message.reply_text(f"🟢 নতুন চ্যানেল সফলভাবে যুক্ত করা হয়েছে: <code>{text}</code>", parse_mode="HTML")
        else:
            await update.effective_message.reply_text("⚠️ এই চ্যানেলটি ইতিমধ্যে লিস্টে আছে।")
        return True

    elif waiting == "rem_channel":
        context.user_data["waiting_for"] = None
        channels = DATA["settings"].setdefault("channels", [])
        if text in channels:
            channels.remove(text)
            save_db()
            await update.effective_message.reply_text(f"🟢 চ্যানেল রিমুভ করা হয়েছে: <code>{text}</code>", parse_mode="HTML")
        else:
            await update.effective_message.reply_text("🔴 এই চ্যানেলটি লিস্টে পাওয়া যায়নি।")
        return True

    elif waiting == "add_admin":
        context.user_data["waiting_for"] = None
        try:
            aid = int(text)
            admin_ids = DATA["settings"].setdefault("admin_ids", [DEFAULT_ADMIN_ID])
            if aid not in admin_ids:
                admin_ids.append(aid)
                save_db()
                await update.effective_message.reply_text(f"👑 নতুন অ্যাডমিন সফলভাবে যুক্ত করা হয়েছে: <code>{aid}</code>", parse_mode="HTML")
            else:
                await update.effective_message.reply_text("⚠️ এই আইডিটি ইতিমধ্যে অ্যাডমিন লিস্টে আছে।")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যাসূচক ইউজার আইডি দিন।")
        return True

    elif waiting == "rem_admin":
        context.user_data["waiting_for"] = None
        try:
            aid = int(text)
            admin_ids = DATA["settings"].setdefault("admin_ids", [DEFAULT_ADMIN_ID])
            if len(admin_ids) <= 1:
                await update.effective_message.reply_text("🚫 কমপক্ষে ১ জন অ্যাডমিন রাখা বাধ্যতামূলক।")
                return True
            if aid in admin_ids:
                admin_ids.remove(aid)
                save_db()
                await update.effective_message.reply_text(f"🟢 অ্যাডমিন লিস্ট থেকে রিমুভ করা হয়েছে: <code>{aid}</code>", parse_mode="HTML")
            else:
                await update.effective_message.reply_text("🔴 এই আইডিটি অ্যাডমিন লিস্টে পাওয়া যায়নি।")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক ইউজার আইডি দিন।")
        return True

    elif waiting == "custom_date_ids":
        context.user_data["waiting_for"] = None
        req_date = text
        day_data = DATA["days"].get(req_date)
        if not day_data or not day_data["ids"]:
            await update.effective_message.reply_text(f"📭 <code>{req_date}</code> তারিখে কোনো আইডি জমা পড়েনি।", parse_mode="HTML")
            return True

        def gen_custom_date_excel():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = req_date
            ws.append(["UID", "Date"])
            for u in day_data["ids"]:
                ws.append([u, req_date])
            bio = io.BytesIO()
            wb.save(bio)
            bio.seek(0)
            return bio

        try:
            excel_bio = await asyncio.to_thread(gen_custom_date_excel)
            await update.effective_message.reply_document(
                document=excel_bio,
                filename=f"all_ids_{req_date}.xlsx",
                caption=f"📋 {req_date} তারিখের সমস্ত জমা দেওয়া আইডির এক্সেল ফাইল (মোট: {len(day_data['ids'])} টি)।"
            )
        except Exception as e:
            await update.effective_message.reply_text(f"🚫 এক্সেল ফাইল তৈরি করা যায়নি: {e}")
        return True

    elif waiting == "user_date_ids_uid":
        target_uid = text
        if not target_uid.isdigit():
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যাসূচক User ID দিন:")
            return True

        context.user_data["target_user_date_uid"] = target_uid
        context.user_data["waiting_for"] = "user_date_ids_date"
        await update.effective_message.reply_text(
            f"📅 ইউজার আইডি <code>{target_uid}</code> সিলেক্ট হয়েছে। এবার কোন তারিখের আইডি দেখতে চান তা YYYY-MM-DD ফরম্যাটে লিখুন (যেমন: <code>2026-10-18</code>):",
            parse_mode="HTML"
        )
        return True

    elif waiting == "user_date_ids_date":
        context.user_data["waiting_for"] = None
        target_uid = context.user_data.get("target_user_date_uid")
        req_date = text

        day_data = DATA["days"].get(req_date)
        if not day_data or target_uid not in day_data["members"]:
            await update.effective_message.reply_text(f"📭 <code>{req_date}</code> তারিখে এই ইউজারের কোনো সাবমিশন পাওয়া যায়নি।", parse_mode="HTML")
            return True

        mem = day_data["members"][target_uid]
        user_ids = mem.get("all_ids", [])
        if not user_ids:
            await update.effective_message.reply_text("🔴 এই ইউজারের ওই তারিখে কোনো আইডি রেকর্ড নেই।")
            return True

        def gen_user_date_excel():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"User {target_uid}"
            ws.append(["UID", "Date", "User ID", "Name"])
            for u in user_ids:
                ws.append([u, req_date, target_uid, mem.get("name", "User")])
            bio = io.BytesIO()
            wb.save(bio)
            bio.seek(0)
            return bio

        try:
            excel_bio = await asyncio.to_thread(gen_user_date_excel)
            await update.effective_message.reply_document(
                document=excel_bio,
                filename=f"user_{target_uid}_{req_date}.xlsx",
                caption=f"📋 ইউজার <a href='tg://user?id={target_uid}'>{html.escape(mem.get('name', 'User'))}</a> এর {req_date} তারিখের আইডি ফাইল (মোট: {len(user_ids)} টি)।",
                parse_mode="HTML"
            )
        except Exception as e:
            await update.effective_message.reply_text(f"🚫 ফাইল তৈরি করা যায়নি: {e}")
        return True

    elif waiting == "broadcast_message":
        context.user_data["waiting_for"] = None
        mode = context.user_data.get("broadcast_mode", "all")
        all_user_ids = set()

        if mode == "unpaid":
            for d_val in DATA["days"].values():
                for uid, mem in d_val["members"].items():
                    if mem["received"] > 0 and get_user_payment_status(uid, today()) == "Unpaid":
                        all_user_ids.add(int(uid))
        else:
            for d_val in DATA["days"].values():
                for uid in d_val["members"].keys():
                    all_user_ids.add(int(uid))

        if not all_user_ids:
            await update.effective_message.reply_text("📭 এই ফিল্টারে কোনো মেম্বার পাওয়া যায়নি।")
            return True

        success_count = 0
        fail_count = 0
        success_users = []
        fail_users = []
        broadcast_text = f"📢 <b>ADMIN ANNOUNCEMENT</b>\n\n{text}"

        for uid in all_user_ids:
            u_name = "User"
            for d_val in DATA["days"].values():
                if str(uid) in d_val["members"]:
                    u_name = d_val["members"][str(uid)]["name"]
                    break

            try:
                await context.bot.send_message(chat_id=uid, text=broadcast_text, parse_mode="HTML")
                success_users.append(f"• 👤 Name: <a href='tg://user?id={uid}'>{html.escape(u_name)}</a>")
                success_count += 1
                await asyncio.sleep(0.05)
            except Exception:
                fail_users.append(f"• 👤 Name: <a href='tg://user?id={uid}'>{html.escape(u_name)}</a>")
                fail_count += 1

        success_list_str = "\n".join(success_users[:25]) if success_users else "কেউ নেই"
        if len(success_users) > 25:
            success_list_str += f"\n... এবং আরও {len(success_users) - 25} জন"

        fail_list_str = "\n".join(fail_users[:25]) if fail_users else "কেউ নেই"
        if len(fail_users) > 25:
            fail_list_str += f"\n... এবং আরও {len(fail_users) - 25} জন"

        await update.effective_message.reply_text(
            f"✅ <b>ব্রডকাস্ট সম্পন্ন হয়েছে!</b>\n\n"
            f"🟢 সফলভাবে পাঠানো হয়েছে: <code>{success_count}</code> জনের কাছে\n"
            f"<b>সফল ব্যবহারকারীগণ:</b>\n{success_list_str}\n\n"
            f"🔴 ব্যর্থ হয়েছে: <code>{fail_count}</code> জন\n"
            f"<b>ব্যর্থ ব্যবহারকারীগণ:</b>\n{fail_list_str}",
            parse_mode="HTML"
        )
        return True

    elif waiting == "reject_bad_ids":
        context.user_data["waiting_for"] = None
        raw_ids = [line.strip() for line in text.splitlines() if line.strip().isdigit()]
        if not raw_ids:
            await update.effective_message.reply_text("🚫 কোনো বৈধ আইডি পাওয়া যায়নি।")
            return True

        user_rejected_map = {}
        found_count = 0

        for d_date, d_data in DATA["days"].items():
            for bad_id in raw_ids:
                for uid, mem in d_data["members"].items():
                    mem.setdefault("rejected", 0)
                    mem.setdefault("rejected_ids", [])

                    if bad_id in mem["all_ids"]:
                        found_count += 1
                        
                        if bad_id in mem["new_ids"]:
                            mem["new_ids"].remove(bad_id)
                            mem["new"] = max(0, mem["new"] - 1)
                            
                            if mem.get("only_uid_count", 0) > 0:
                                mem["only_uid_count"] -= 1
                            elif mem.get("page_uid_count", 0) > 0:
                                mem["page_uid_count"] -= 1

                        elif bad_id in mem["duplicate_ids"]:
                            mem["duplicate_ids"].remove(bad_id)
                            mem["duplicate"] = max(0, mem["duplicate"] - 1)

                        if bad_id not in mem["rejected_ids"]:
                            mem["rejected"] += 1
                            mem["rejected_ids"].append(bad_id)

                        user_rejected_map.setdefault(uid, []).append(bad_id)

        save_db()

        member_blocks = []
        for target_uid, bad_list in user_rejected_map.items():
            mem_name = DATA["days"][today()]["members"].get(target_uid, {}).get("name", "Member")
            
            tree_ids = []
            for idx, b_id in enumerate(bad_list):
                prefix = "└" if idx == len(bad_list) - 1 else "├"
                tree_ids.append(f"{prefix} 🆔 <code>{b_id}</code>")

            tree_str = "\n".join(tree_ids)
            member_blocks.append(f"👤 Name: <a href='tg://user?id={target_uid}'>{html.escape(mem_name)}</a>\n{tree_str}")

            try:
                formatted_bad_list = "\n".join([f"<code>{i}</code>" for i in bad_list])
                await context.bot.send_message(
                    chat_id=int(target_uid),
                    text=(
                        f"🚫 <b>আপনার পাঠানো নিচের <code>{len(bad_list)}</code> টি আইডি রিজেক্ট করা হয়েছে:</b>\n\n"
                        + formatted_bad_list +
                        "\n\n⚠️ আপনার ব্যালেন্স ও স্ট্যাটাস থেকে এটার মূল্য মাইনাস করা হয়েছে।"
                    ),
                    parse_mode="HTML"
                )
            except Exception as e:
                print(f"User Notification Error ({target_uid}):", e)

        details_msg = "\n\n".join(member_blocks) if member_blocks else "কোনো মেম্বারের ডাটাবেজে রেকর্ড পাওয়া যায়নি।"

        await update.effective_message.reply_text(
            f"✅ <b>প্রসেসিং সম্পন্ন হয়েছে!</b>\n\n"
            f"📥 মোট সাবমিট রিজেক্ট আইডি: <code>{len(raw_ids)}</code> টি\n"
            f"🔍 ডাটাবেজে পাওয়া গেছে: <code>{found_count}</code> টি\n\n"
            f"<b>📋 মেম্বার ভিত্তিক রিজেক্টের বিবরণ:</b>\n\n{details_msg}",
            parse_mode="HTML"
        )
        return True

    elif waiting == "set_price_only":
        context.user_data["waiting_for"] = None
        try:
            val = float(text)
            DATA["settings"]["price_only_uid"] = val
            save_db()
            await update.effective_message.reply_text(f"🟢 <b>ONLY UID</b> এর প্রাইস সেট করা হয়েছে: <code>{val} BDT</code>", parse_mode="HTML")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যা বা দশমিক লিখুন।")
        return True

    elif waiting == "set_price_page":
        context.user_data["waiting_for"] = None
        try:
            val = float(text)
            DATA["settings"]["price_page_uid"] = val
            save_db()
            await update.effective_message.reply_text(f"🟢 <b>PAGE UID</b> এর প্রাইস সেট করা হয়েছে: <code>{val} BDT</code>", parse_mode="HTML")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যা বা দশমিক লিখুন।")
        return True

    elif waiting == "set_work_schedule_time":
        context.user_data["waiting_for"] = None
        parts = text.split("-")
        if len(parts) == 2:
            DATA["settings"]["start_time"] = parts[0].strip()
            DATA["settings"]["end_time"] = parts[1].strip()
            save_db()
            await update.effective_message.reply_text(f"🟢 কাজের সময়সীমা আপডেট করা হয়েছে: <code>{parts[0].strip()} - {parts[1].strip()}</code>", parse_mode="HTML")
        else:
            await update.effective_message.reply_text("🔴 সঠিক ফরম্যাটে সময় লিখুন (যেমন: 10:00-22:00)।")
        return True

    elif waiting == "set_pay_template":
        context.user_data["waiting_for"] = None
        DATA["settings"]["payment_msg_template"] = text
        save_db()
        await update.effective_message.reply_text("🟢 পেমেন্ট মেসেজ টেমপ্লেট আপডেট করা হয়েছে।", parse_mode="HTML")
        return True

    elif waiting == "view_custom_date":
        context.user_data["waiting_for"] = None
        req_date = text
        day_data = DATA["days"].get(req_date)
        if not day_data or not day_data["members"]:
            await update.effective_message.reply_text(f"📭 <code>{req_date}</code> তারিখে কোনো সাবমিশন পাওয়া যায়নি।", parse_mode="HTML")
            return True

        lines = [f"<b>📊 SUBMISSION REPORT FOR: {req_date}</b>\n"]
        total_r = total_n = total_d = total_rej = 0
        for uid, mem in day_data["members"].items():
            if mem["received"] > 0:
                earned = calculate_member_earned(mem)
                pay_st = get_user_payment_status(uid, req_date)
                total_r += mem["received"]
                total_n += mem["new"]
                total_d += mem["duplicate"]
                rej_c = mem.get("rejected", 0)
                total_rej += rej_c
                o_cnt = mem.get("only_uid_count", 0)
                p_cnt = mem.get("page_uid_count", 0)

                lines.extend([
                    f"👤 Name: <a href='tg://user?id={uid}'>{html.escape(mem['name'])}</a>",
                    f"📥 Total Received: <code>{mem['received']}</code>",
                    f"🟢 New IDs: <code>{mem['new']}</code> (Only: {o_cnt} | Page: {p_cnt})",
                    f"♻️ Duplicate IDs: <code>{mem['duplicate']}</code>",
                    f"🚫 Rejected IDs: <code>{rej_c}</code>",
                    f"💰 Earned Amount: <code>{earned:.2f} BDT</code> ({pay_st})",
                    "────────────────────\n"
                ])

        lines.extend([
            "━━━━━━━━━━━━━━━━━━",
            f"📦 <b>Total Received:</b> <code>{total_r}</code> | 🟢 <b>New:</b> <code>{total_n}</code>"
        ])
        
        full_text = "\n".join(lines)
        for start_pos in range(0, len(full_text), 3500):
            await update.effective_message.reply_text(full_text[start_pos:start_pos + 3500], parse_mode="HTML")
        return True

    elif waiting == "export_custom_date":
        context.user_data["waiting_for"] = None
        req_date = text
        day_data = DATA["days"].get(req_date)
        if not day_data or not day_data["members"]:
            await update.effective_message.reply_text(f"📭 <code>{req_date}</code> তারিখে কোনো ডাটা নেই।", parse_mode="HTML")
            return True

        def gen_excel():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = req_date
            ws.append(["Date", "User ID", "Name", "Received", "New", "Only UID", "Page UID", "Duplicate", "Rejected", "Earned BDT", "Payment Info", "Payment Status"])
            for u_id, mem in day_data["members"].items():
                if mem["received"] > 0:
                    earned = calculate_member_earned(mem)
                    pay_st = get_user_payment_status(u_id, req_date)
                    pay_dets = DATA["user_payment_details"].get(str(u_id), {})
                    pay_str = ", ".join([f"{m}: {n}" for m, n in pay_dets.items()]) if isinstance(pay_dets, dict) and pay_dets else "Not Set"
                    ws.append([req_date, u_id, mem["name"], mem["received"], mem["new"], mem.get("only_uid_count",0), mem.get("page_uid_count",0), mem["duplicate"], mem.get("rejected", 0), earned, pay_str, pay_st])
            bio = io.BytesIO()
            wb.save(bio)
            bio.seek(0)
            return bio

        excel_bio = await asyncio.to_thread(gen_excel)
        await update.effective_message.reply_document(
            document=excel_bio,
            filename=f"report_{req_date}.xlsx",
            caption=f"📊 Report for {req_date}"
        )
        return True

    elif waiting == "view_user_hist":
        context.user_data["waiting_for"] = None
        target_uid = text
        dates = sorted(DATA["days"].keys(), reverse=True)
        lines = [f"👤 Name: <b>USER HISTORY FOR:</b> <code>{target_uid}</code>\n"]
        found = False

        for d_date in dates:
            mem = DATA["days"][d_date]["members"].get(target_uid)
            if mem:
                found = True
                rej_c = mem.get("rejected", 0)
                earned = calculate_member_earned(mem)
                pay_st = get_user_payment_status(target_uid, d_date)
                lines.append(f"📅 <code>{d_date}</code>\n  📥 Total: <code>{mem['received']}</code> | 🟢 New: <code>{mem['new']}</code> (Only: {mem.get('only_uid_count',0)}, Page: {mem.get('page_uid_count',0)})\n  ♻️ Dup: <code>{mem['duplicate']}</code> | 🚫 Rej: <code>{rej_c}</code>\n  💰 Earned: <code>{earned:.2f} BDT</code> ({pay_st})\n")

        if found:
            await update.effective_message.reply_text("\n".join(lines), parse_mode="HTML")
        else:
            await update.effective_message.reply_text("🔴 এই ইউজার আইডির কোনো রেকর্ড পাওয়া যায়নি।")
        return True

    elif waiting == "del_hist_type":
        context.user_data["waiting_for"] = None
        input_val = text
        if input_val in DATA["days"]:
            del DATA["days"][input_val]
            save_db()
            await update.effective_message.reply_text(f"🗑️ <code>{input_val}</code> তারিখের সমস্ত হিস্ট্রি সফলভাবে ডিলিট করা হয়েছে।", parse_mode="HTML")
        else:
            deleted = False
            for d_date in DATA["days"]:
                if input_val in DATA["days"][d_date]["members"]:
                    del DATA["days"][d_date]["members"][input_val]
                    deleted = True
            if deleted:
                save_db()
                await update.effective_message.reply_text(f"🗑️ ইউজার <code>{input_val}</code> এর সমস্ত হিস্ট্রি ডিলিট করা হয়েছে।", parse_mode="HTML")
            else:
                await update.effective_message.reply_text("🔴 প্রদত্ত তারিখ বা ইউজার আইডি ডাটাবেজে পাওয়া যায়নি।")
        return True

    elif waiting == "add_rcv":
        context.user_data["waiting_for"] = None
        try:
            cid = int(text)
            if cid not in DATA["settings"]["target_chat_ids"]:
                DATA["settings"]["target_chat_ids"].append(cid)
                save_db()
                await update.effective_message.reply_text(f"🟢 সফলভাবে রিসিভার চ্যাট আইডি যোগ করা হয়েছে: <code>{cid}</code>", parse_mode="HTML")
            else:
                await update.effective_message.reply_text("⚠️ এই আইডিটি ইতিমধ্যে লিস্টে আছে।")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যাসূচক চ্যাট আইডি দিন।")
        return True

    elif waiting == "rem_rcv":
        context.user_data["waiting_for"] = None
        try:
            cid = int(text)
            if cid in DATA["settings"]["target_chat_ids"]:
                DATA["settings"]["target_chat_ids"].remove(cid)
                save_db()
                await update.effective_message.reply_text(f"🟢 রিসিভার আইডি রিমুভ করা হয়েছে: <code>{cid}</code>", parse_mode="HTML")
            else:
                await update.effective_message.reply_text("🔴 এই আইডিটি লিস্টে পাওয়া যায়নি।")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক আইডি দিন।")
        return True

    elif waiting == "set_limit":
        context.user_data["waiting_for"] = None
        try:
            limit = int(text)
            DATA["settings"]["min_ids_limit"] = limit
            save_db()
            await update.effective_message.reply_text(f"🟢 মিনিমাম আইডি লিমিট সেট করা হয়েছে: <code>{limit}</code> টি", parse_mode="HTML")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যা দিন।")
        return True

    elif waiting == "ban_user":
        context.user_data["waiting_for"] = None
        try:
            bid = int(text)
            if bid in DATA["banned_users"]:
                DATA["banned_users"].remove(bid)
                save_db()
                await update.effective_message.reply_text(f"🟢 ইউজারকে আনব্যান করা হয়েছে: <code>{bid}</code>", parse_mode="HTML")
            else:
                DATA["banned_users"].append(bid)
                save_db()
                await update.effective_message.reply_text(f"🚫 ইউজারকে ব্যান করা হয়েছে: <code>{bid}</code>", parse_mode="HTML")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক ইউজার আইডি দিন।")
        return True

    return False


async def check_payment_set(user, update_msg):
    uid_str = str(user.id)
    pay_info = DATA["user_payment_details"].get(uid_str)
    if not pay_info or not isinstance(pay_info, dict) or not any(pay_info.values()):
        await update_msg.reply_text(
            "⚠️ <b>পেমেন্ট নম্বর সেট করা হয়নি!</b>\n\n"
            "আইডি জমা দেওয়ার আগে দয়া করে আপনার পেমেন্ট নম্বর সেট করুন。\n"
            "👉 পেমেন্ট নম্বর সেট করতে <b>/setpayment</b> কমান্ড দিন।",
            parse_mode="HTML"
        )
        return False
    return True


async def receive_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    
    if update.effective_message.document and is_bot_admin(user.id):
        doc = update.effective_message.document
        if doc.file_name and doc.file_name.endswith(".json"):
            file = await context.bot.get_file(doc.file_id)
            bio = io.BytesIO()
            await file.download_to_memory(bio)
            bio.seek(0)
            try:
                new_data = json.loads(bio.read().decode("utf-8"))
                if isinstance(new_data, dict) and "settings" in new_data:
                    global DATA
                    DATA = new_data
                    if "Binance" not in DATA.get("payment_methods", []):
                        DATA.setdefault("payment_methods", ["bKash", "Nagad", "Rocket", "Binance"])
                    save_db()
                    await update.effective_message.reply_text("✅ <b>ডাটাবেজ ব্যাকআপ ফাইল সফলভাবে রিস্টোর (Restore) করা হয়েছে!</b>", parse_mode="HTML")
                    return
            except Exception as e:
                await update.effective_message.reply_text(f"🚫 ডাটাবেজ রিস্টোর ব্যর্থ হয়েছে: {e}")
                return

    if user.id in DATA["banned_users"]:
        return
    if DATA["settings"]["maintenance"] and not is_bot_admin(user.id):
        await update.effective_message.reply_text("🛠️ বট বর্তমানে মেনটেইনেন্স মোডে আছে।")
        return
    if not await is_joined(user.id, context):
        await require_join(update, context)
        return

    text = update.effective_message.text or ""

    if is_bot_admin(user.id) and text == "👑 অ্যাডমিন প্যানেল":
        await admin(update, context)
        return

    if context.user_data.get("waiting_for"):
        handled = await handle_admin_inputs(update, context)
        if handled:
            return

    ids = [line.strip() for line in text.splitlines() if line.strip().isdigit()]
    if not ids:
        return

    if not is_bot_admin(user.id):
        if not await check_payment_set(user, update.effective_message):
            return

        if not is_work_time():
            st = DATA['settings'].get('start_time', '10:00')
            et = DATA['settings'].get('end_time', '22:00')
            await update.effective_message.reply_text(
                f"🚫 <b>কাজের সময়সূচী বন্ধ রয়েছে!</b>\n\n⏰ <b>কাজের সময়:</b> {st} থেকে {et} পর্যন্ত。\n⌛ নির্ধারিত সময়ে আইডি জমা দিন।",
                parse_mode="HTML"
            )
            return

    await ask_category_selection(update, ids)


async def ask_category_selection(update: Update, ids, file_info=None):
    sub_temp_id = f"temp_{update.effective_user.id}_{int(datetime.now().timestamp())}"
    context_map = getattr(ask_category_selection, "temp_submissions", {})
    context_map[sub_temp_id] = {"ids": ids, "file_info": file_info}
    ask_category_selection.temp_submissions = context_map

    keyboard = [
        [
            InlineKeyboardButton("🆔 ONLY UID", callback_data=f"cat_sel_ONLY_{sub_temp_id}"),
            InlineKeyboardButton("📄 PAGE UID", callback_data=f"cat_sel_PAGE_{sub_temp_id}")
        ]
    ]

    await update.effective_message.reply_text(
        f"🏷️ আপনি <code>{len(ids)}</code> টি আইডি পাঠিয়েছেন। অনুগ্রহ করে ক্যাটাগরি সিলেক্ট করুন:",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="HTML"
    )


# ================= CLEAN SUB-MENU ADMIN PANEL =================

async def admin(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    context.user_data.clear()

    if not is_bot_admin(user.id):
        await update.effective_message.reply_text("🚫 আপনার এই কমান্ডটি ব্যবহার করার অনুমতি নেই।")
        return

    keyboard = [
        [InlineKeyboardButton("📊 Reports & Summary", callback_data="cat_reports")],
        [InlineKeyboardButton("💵 Price & Payment Control", callback_data="cat_pricing")],
        [InlineKeyboardButton("🛠️ ID & Data Management", callback_data="cat_id_mgt")],
        [InlineKeyboardButton("⚙️ Bot System Settings", callback_data="cat_sys_settings")],
        [InlineKeyboardButton("👥 User & Receiver Control", callback_data="cat_user_control")]
    ]

    text = "👑 <b>MAIN ADMIN CONTROL PANEL</b>\n\nযে ক্যাটাগরির কাজ করতে চান নিচের বাটন থেকে সিলেক্ট করুন:"
    
    if update.callback_query:
        await update.callback_query.edit_message_text(text, reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
    else:
        await update.effective_message.reply_text(text, reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")


async def render_sys_settings(query):
    s = DATA["settings"]
    status_icon = "🔴 OFF" if s["maintenance"] else "🟢 ON"
    approval_icon = "🟢 ON" if s["approval_mode"] else "⚡ Direct"
    mode_icon = "🔄 All-Time" if s["duplicate_mode"] == "all_time" else "📅 Daily"
    sch_icon = "🟢 ON" if s.get("work_schedule_enabled") else "🔴 OFF"
    file_toggle_icon = "🟢 ON" if s.get("send_live_check_file", True) else "🔴 OFF"

    keyboard = [
        [InlineKeyboardButton(f"🛡️ Approval Mode: {approval_icon}", callback_data="adm_toggle_approval")],
        [InlineKeyboardButton(f"🛠️ Bot Status: {status_icon}", callback_data="adm_toggle_maint")],
        [InlineKeyboardButton(f"⚙️ Dup Check Mode: {mode_icon}", callback_data="adm_toggle_mode")],
        [InlineKeyboardButton(f"⏰ Work Schedule: {sch_icon}", callback_data="adm_toggle_schedule")],
        [InlineKeyboardButton(f"📁 Live Check File: {file_toggle_icon}", callback_data="adm_toggle_live_file")],
        [InlineKeyboardButton("🕒 Set Work Hours", callback_data="adm_set_schedule_time")],
        [InlineKeyboardButton("🔢 Min Limit", callback_data="adm_limit"), InlineKeyboardButton("📢 Manage Channels", callback_data="adm_channels_menu")],
        [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
    ]
    await query.edit_message_text("⚙️ <b>BOT SYSTEM SETTINGS</b>\n\nবটের গ্লোবাল সেটিংস পরিবর্তন করুন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")


async def render_unpaid_list(query_or_msg):
    unpaid_buttons = []
    for d_date, d_val in DATA["days"].items():
        for uid, mem in d_val["members"].items():
            if mem["received"] > 0 and get_user_payment_status(uid, d_date) == "Unpaid":
                btn_text = f"👤 Name: {mem['name']} | 📅 {d_date} | 💰 {calculate_member_earned(mem):.2f} BDT"
                unpaid_buttons.append([InlineKeyboardButton(btn_text, callback_data=f"paydet_{uid}_{d_date}")])

    if not unpaid_buttons:
        txt = "🎉 বর্তমানে কোনো Unpaid পেমেন্ট বাকি নেই!"
        if hasattr(query_or_msg, 'edit_message_text'):
            await query_or_msg.edit_message_text(txt, parse_mode="HTML")
        else:
            await query_or_msg.reply_text(txt, parse_mode="HTML")
        return

    unpaid_buttons.append([InlineKeyboardButton("🔙 Back", callback_data="cat_pricing")])
    txt = "💳 <b>SMART UNPAID PAYMENT LIST</b>\n\nযে ইউজারের পেমেন্ট ডিটেইলস দেখতে চান তার ওপর ক্লিক করুন:"
    if hasattr(query_or_msg, 'edit_message_text'):
        await query_or_msg.edit_message_text(txt, reply_markup=InlineKeyboardMarkup(unpaid_buttons), parse_mode="HTML")
    else:
        await query_or_msg.reply_text(txt, reply_markup=InlineKeyboardMarkup(unpaid_buttons), parse_mode="HTML")


async def admin_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    user = update.effective_user

    try:
        await query.answer()
    except Exception:
        pass

    data = query.data

    try:
        if data == "member_custom_date_ids":
            context.user_data["waiting_for"] = "member_enter_date"
            await query.message.reply_text(
                "📅 আপনি কোন তারিখের আইডি দেখতে চান তা YYYY-MM-DD ফরম্যাটে লিখুন (যেমন: <code>2026-10-18</code>):",
                parse_mode="HTML"
            )
            return

        if data.startswith("setlang_"):
            lang_code = data.replace("setlang_", "")
            DATA["user_languages"][str(user.id)] = lang_code
            save_db()
            msg_txt = "✅ ভাষা সফলভাবে পরিবর্তন করা হয়েছে!" if lang_code == "bn" else "✅ Language updated successfully!"
            await query.edit_message_text(msg_txt)
            return

        if data.startswith("setpaym_"):
            method = data.replace("setpaym_", "")
            context.user_data["selected_pay_method"] = method
            context.user_data["waiting_for"] = "user_enter_pay_num"
            await query.edit_message_text(
                f"💳 <b>{method}</b> সিলেক্ট করেছেন। এবার আপনার {method} নম্বরটি লিখে পাঠান:",
                parse_mode="HTML"
            )
            return

        if data.startswith("cat_sel_"):
            parts = data.split("_", 3)
            cat_type = parts[2]
            sub_temp_id = parts[3]
            
            category_name = "ONLY UID" if cat_type == "ONLY" else "PAGE UID"

            context_map = getattr(ask_category_selection, "temp_submissions", {})
            temp_data = context_map.get(sub_temp_id)

            if not temp_data:
                await query.edit_message_text("⚠️ অনুরোধের সময়সীমা পার হয়ে গেছে। পুনরায় পাঠাক।")
                return

            del context_map[sub_temp_id]
            
            await query.edit_message_text(f"✅ Selected Category: <b>{category_name}</b>", parse_mode="HTML")
            await handle_incoming_data(user, temp_data["ids"], category_name, context, query.message, temp_data["file_info"])
            return

        if not is_bot_admin(user.id):
            return

        context.user_data.clear()

        if data == "cat_main":
            await admin(update, context)
            return

        if data == "cat_reports":
            keyboard = [
                [InlineKeyboardButton("👥 Today Member Summary", callback_data="adm_member_summary")],
                [InlineKeyboardButton("📋 Today All Received IDs", callback_data="adm_all_ids_filter")],
                [InlineKeyboardButton("📅 View All IDs for Custom Date", callback_data="adm_custom_date_ids_btn")],
                [InlineKeyboardButton("👤 View User's IDs by Date", callback_data="adm_user_date_ids_btn")],
                [InlineKeyboardButton("📅 View Custom Date Report", callback_data="adm_view_custom_date")],
                [InlineKeyboardButton("👤 View Specific User History", callback_data="adm_user_hist")],
                [InlineKeyboardButton("👥 All Registered Members", callback_data="adm_all_members")],
                [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
            ]
            await query.edit_message_text("📊 <b>REPORTS & SUMMARY MENU</b>\n\nপ্রয়োজনীয় রিপোর্ট বেছে নিন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
            return

        if data == "adm_custom_date_ids_btn":
            context.user_data["waiting_for"] = "custom_date_ids"
            await query.message.reply_text("📅 যে তারিখের সমস্ত আইডি দেখতে চান তা YYYY-MM-DD ফরম্যাটে লিখুন (যেমন: <code>2026-10-18</code>):", parse_mode="HTML")
            return

        if data == "adm_user_date_ids_btn":
            context.user_data["waiting_for"] = "user_date_ids_uid"
            await query.message.reply_text("👤 যে মেম্বারের আইডি দেখতে চান তার <b>User ID</b> লিখে পাঠান:", parse_mode="HTML")
            return

        if data == "cat_pricing":
            p_only = DATA["settings"].get("price_only_uid", 0.0)
            p_page = DATA["settings"].get("price_page_uid", 0.0)
            keyboard = [
                [InlineKeyboardButton(f"✏️ Set ONLY UID Price (Cur: {p_only} BDT)", callback_data="adm_set_p_only")],
                [InlineKeyboardButton(f"✏️ Set PAGE UID Price (Cur: {p_page} BDT)", callback_data="adm_set_p_page")],
                [InlineKeyboardButton("💳 Smart Unpaid Payment List", callback_data="adm_smart_unpaid")],
                [InlineKeyboardButton("✉️ Edit Payment Success Msg", callback_data="adm_edit_pay_msg")],
                [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
            ]
            await query.edit_message_text("💵 <b>PRICE & PAYMENT CONTROL</b>\n\nদাম ও পেমেন্ট আপডেট করুন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
            return

        if data == "cat_id_mgt":
            keyboard = [
                [InlineKeyboardButton("🚫 Reject Specific Bad IDs", callback_data="adm_reject_bad_ids")],
                [InlineKeyboardButton("🗑️ Delete Single Date/User Data", callback_data="adm_del_hist")],
                [InlineKeyboardButton("🔥 Reset All History Data", callback_data="adm_reset_all")],
                [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
            ]
            await query.edit_message_text("🛠️ <b>ID & DATA MANAGEMENT</b>\n\nআইডি ও ডাটা সংক্রান্ত কাজ করুন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
            return

        if data == "cat_sys_settings":
            await render_sys_settings(query)
            return

        if data == "adm_toggle_live_file":
            DATA["settings"]["send_live_check_file"] = not DATA["settings"].get("send_live_check_file", True)
            save_db()
            await render_sys_settings(query)
            return

        if data == "adm_channels_menu":
            channels_list = "\n".join([f"• <code>{ch}</code>" for ch in DATA["settings"].get("channels", [DEFAULT_CHANNEL])])
            keyboard = [
                [InlineKeyboardButton("➕ Add Channel", callback_data="adm_add_channel"), InlineKeyboardButton("➖ Remove Channel", callback_data="adm_rem_channel")],
                [InlineKeyboardButton("🔙 Back to Settings", callback_data="cat_sys_settings")]
            ]
            await query.edit_message_text(
                f"📢 <b>MANAGE JOIN CHANNELS</b>\n\nবর্তমান চ্যানেলসমূহ:\n{channels_list}\n\nচ্যানেল যোগ বা বাদ দিতে নিচের বাটন চাপুন:",
                reply_markup=InlineKeyboardMarkup(keyboard),
                parse_mode="HTML"
            )
            return

        if data == "adm_add_channel":
            context.user_data["waiting_for"] = "add_channel"
            await query.message.reply_text("📢 যে চ্যানেলটি যুক্ত করতে চান তার ইউজারনেম বা লিংক পাঠান (যেমন: `@OMORTECH`):", parse_mode="HTML")
            return

        if data == "adm_rem_channel":
            context.user_data["waiting_for"] = "rem_channel"
            await query.message.reply_text("📢 যে চ্যানেলটি রিমুভ করতে চান তার ইউজারনেম বা লিংক লিখে পাঠান:", parse_mode="HTML")
            return

        if data == "cat_user_control":
            keyboard = [
                [InlineKeyboardButton("➕ Add Receiver", callback_data="adm_add_rcv"), InlineKeyboardButton("➖ Remove Receiver", callback_data="adm_rem_rcv")],
                [InlineKeyboardButton("👑 Manage Admins", callback_data="adm_admins_menu")],
                [InlineKeyboardButton("📢 Broadcast (All Users)", callback_data="adm_broadcast_all")],
                [InlineKeyboardButton("📢 Broadcast (Unpaid Only)", callback_data="adm_broadcast_unpaid")],
                [InlineKeyboardButton("✉️ Message Specific User", callback_data="adm_msg_user")],
                [InlineKeyboardButton("🚫 Ban / Unban User", callback_data="adm_ban")],
                [InlineKeyboardButton("💾 Backup Database File", callback_data="adm_backup_db")],
                [InlineKeyboardButton("📥 Export Full Excel Report", callback_data="adm_export")],
                [InlineKeyboardButton("📅 Export Specific Date Excel", callback_data="adm_export_custom_date_btn")],
                [InlineKeyboardButton("📋 View Full System Config", callback_data="adm_view")],
                [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
            ]
            await query.edit_message_text("👥 <b>USER & RECEIVER CONTROL</b>\n\nরিসিভার চ্যানেল বা ইউজার নিয়ন্ত্রণ করুন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
            return

        if data == "adm_admins_menu":
            admins_list = "\n".join([f"• <code>{aid}</code>" for aid in DATA["settings"].get("admin_ids", [DEFAULT_ADMIN_ID])])
            keyboard = [
                [InlineKeyboardButton("➕ Add Admin", callback_data="adm_add_admin"), InlineKeyboardButton("➖ Remove Admin", callback_data="adm_rem_admin")],
                [InlineKeyboardButton("🔙 Back to User Control", callback_data="cat_user_control")]
            ]
            await query.edit_message_text(
                f"👑 <b>MANAGE BOT ADMINS</b>\n\nবর্তমান অ্যাডমিনগণ:\n{admins_list}\n\nঅ্যাডমিন যোগ বা বাদ দিতে নিচের বাটন চাপুন:",
                reply_markup=InlineKeyboardMarkup(keyboard),
                parse_mode="HTML"
            )
            return

        if data == "adm_add_admin":
            context.user_data["waiting_for"] = "add_admin"
            await query.message.reply_text("👑 নতুন অ্যাডমিনের <b>User ID</b> লিখে পাঠান:", parse_mode="HTML")
            return

        if data == "adm_rem_admin":
            context.user_data["waiting_for"] = "rem_admin"
            await query.message.reply_text("👑 যে অ্যাডমিনকে রিমুভ করতে চান তার <b>User ID</b> লিখে পাঠান:", parse_mode="HTML")
            return

        if data == "adm_backup_db":
            try:
                with open(DB_FILE, "rb") as f:
                    await query.message.reply_document(
                        document=f,
                        filename=f"bot_data_backup_{today()}.json",
                        caption="💾 <b>আপনার বটের সম্পূর্ণ ডাটাবেজ ব্যাকআপ ফাইল।</b>",
                        parse_mode="HTML"
                    )
            except Exception as e:
                await query.message.reply_text(f"🚫 ব্যাকআপ নিতে সমস্যা হয়েছে: {e}")
            return

        if data == "adm_toggle_schedule":
            DATA["settings"]["work_schedule_enabled"] = not DATA["settings"].get("work_schedule_enabled", False)
            save_db()
            await render_sys_settings(query)
            return

        if data == "adm_limit":
            context.user_data["waiting_for"] = "set_limit"
            await query.message.reply_text(
                f"🔢 মিনিমাম আইডির সংখ্যা লিখুন। বর্তমান: <code>{DATA['settings']['min_ids_limit']}</code>", 
                parse_mode="HTML"
            )
            return

        if data == "adm_set_schedule_time":
            context.user_data["waiting_for"] = "set_work_schedule_time"
            cur_s = DATA["settings"].get("start_time", "10:00")
            cur_e = DATA["settings"].get("end_time", "22:00")
            await query.message.reply_text(
                f"🕒 <b>SET WORK HOURS</b>\n\nবর্তমান সময়: <code>{cur_s} - {cur_e}</code>\n\n"
                "নতুন সময়সীমা HH:MM-HH:MM ফরম্যাটে লিখে পাঠান (যেমন: <code>10:00-22:00</code>):",
                parse_mode="HTML"
            )
            return

        if data == "adm_broadcast_all":
            context.user_data["waiting_for"] = "broadcast_message"
            context.user_data["broadcast_mode"] = "all"
            await query.message.reply_text("📢 <b>BROADCAST TO ALL MEMBERS</b>\n\nসকল রেজিস্টার্ড মেম্বারদের জন্য মেসেজটি লিখে পাঠান:", parse_mode="HTML")
            return

        if data == "adm_broadcast_unpaid":
            context.user_data["waiting_for"] = "broadcast_message"
            context.user_data["broadcast_mode"] = "unpaid"
            await query.message.reply_text("📢 <b>BROADCAST TO UNPAID MEMBERS ONLY</b>\n\nযাদের পেমেন্ট বাকি আছে তাদের জন্য মেসেজটি লিখে পাঠান:", parse_mode="HTML")
            return

        if data == "adm_smart_unpaid":
            await render_unpaid_list(query)
            return

        if data.startswith("paydet_"):
            _, p_uid, p_date = data.split("_", 2)
            day_data = DATA["days"].get(p_date)
            if not day_data or p_uid not in day_data["members"]:
                await query.message.reply_text("🔴 মেম্বারের ডাটা পাওয়া যায়নি।")
                return

            mem = day_data["members"][p_uid]
            earned = calculate_member_earned(mem)
            
            pay_details = DATA["user_payment_details"].get(p_uid, {})
            if isinstance(pay_details, dict) and pay_details:
                pay_lines = [f"  • {m}: <code>{n}</code>" for m, n in pay_details.items() if n]
                pay_info = "\n" + "\n".join(pay_lines) if pay_lines else " <code>নম্বর সেট করা হয়নি</code>"
            else:
                pay_info = " <code>নম্বর সেট করা হয়নি</code>"

            rej_c = mem.get("rejected", 0)
            only_c = mem.get("only_uid_count", 0)
            page_c = mem.get("page_uid_count", 0)

            detail_text = (
                f"👤 Name: {html.escape(mem['name'])}\n"
                f"🆔 <b>User ID:</b> <code>{p_uid}</code>\n"
                f"📅 <b>Date:</b> <code>{p_date}</code>\n"
                f"💳 <b>Payment Number:</b>{pay_info}\n"
                f"━━━━━━━━━━━━━━━━━━\n"
                f"📥 Total Received: <code>{mem['received']}</code>\n"
                f"🟢 New: <code>{mem['new']}</code> (Only: {only_c}, Page: {page_c})\n"
                f"♻️ Duplicate: <code>{mem['duplicate']}</code>\n"
                f"🚫 Rejected: <code>{rej_c}</code>\n"
                f"💰 <b>Net Earned:</b> <code>{earned:.2f} BDT</code>\n"
                f"💳 <b>Status:</b> Unpaid"
            )

            keyboard = [
                [InlineKeyboardButton("✅ Mark as Paid", callback_data=f"do_paid_{p_uid}_{p_date}")],
                [InlineKeyboardButton("🔙 Back to Unpaid List", callback_data="adm_smart_unpaid")]
            ]
            await query.edit_message_text(detail_text, reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
            return

        if data.startswith("do_paid_"):
            _, _, p_uid, p_date = data.split("_", 3)
            day_data = DATA["days"].get(p_date)
            if day_data and p_uid in day_data["members"]:
                mem = day_data["members"][p_uid]
                earned = calculate_member_earned(mem)
                DATA["payments"][f"{p_uid}_{p_date}"] = "Paid"
                save_db()

                template = DATA["settings"].get("payment_msg_template", "🎉 <b>আপনার পেমেন্ট সফলভাবে পরিশোধ করা হয়েছে!</b>\n\n💵 <b>পরিশোধিত অর্থ:</b> {amount} BDT")
                msg_to_send = template.replace("{amount}", f"{earned:.2f}").replace("{date}", p_date)

                try:
                    await context.bot.send_message(chat_id=int(p_uid), text=msg_to_send, parse_mode="HTML")
                except Exception as e:
                    print("Payment Notif Failed:", e)

                await query.answer("✅ পেমেন্ট সফলভাবে PAID হিসেবে মার্ক করা হয়েছে!", show_alert=True)
                await render_unpaid_list(query)
            return

        if data == "adm_member_summary":
            date = today()
            day_data = DATA["days"].get(date)
            
            if not day_data or not day_data["members"]:
                await query.message.reply_text(f"📅 <b>Date:</b> <code>{date}</code>\n\n📭 আজকের দিনে এখনো কোনো আইডি জমা পড়েনি।", parse_mode="HTML")
                return

            blocks = [
                "<b>📊 TODAY'S MEMBER SUMMARY</b>",
                f"📅 <b>Date:</b> <code>{date}</code>",
                "━━━━━━━━━━━━━━━━━━━━\n"
            ]

            total_day_received = total_day_new = total_day_dup = total_day_rej = 0
            active_count = 0

            for uid, mem in day_data["members"].items():
                if mem['received'] > 0:
                    active_count += 1
                    adm_rep_name = f"<a href='tg://user?id={uid}'>{html.escape(mem['name'])}</a>"
                    total_day_received += mem['received']
                    total_day_new += mem['new']
                    total_day_dup += mem['duplicate']
                    rej_c = mem.get("rejected", 0)
                    total_day_rej += rej_c
                    earned = calculate_member_earned(mem)
                    pay_st = get_user_payment_status(uid, date)

                    o_cnt = mem.get("only_uid_count", 0)
                    p_cnt = mem.get("page_uid_count", 0)

                    m_block = (
                        f"👤 Name: {adm_rep_name}\n"
                        f"📥 Total Received: <code>{mem['received']}</code>\n"
                        f"🟢 New IDs: <code>{mem['new']}</code> (Only: {o_cnt} | Page: {p_cnt})\n"
                        f"♻️ Duplicate IDs: <code>{mem['duplicate']}</code>\n"
                        f"🚫 Rejected IDs: <code>{rej_c}</code>\n"
                        f"💰 Earned Amount: <code>{earned:.2f} BDT</code> ({pay_st})\n"
                        f"────────────────────"
                    )
                    blocks.append(m_block)

            if active_count == 0:
                await query.message.reply_text(f"📅 <b>Date:</b> <code>{date}</code>\n\n📭 আজকের দিনে কেউ কোনো আইডি জমা দেয়নি।", parse_mode="HTML")
                return

            blocks.append(
                "━━━━━━━━━━━━━━━━━━━━\n"
                f"📦 <b>Today's Total Received:</b> <code>{total_day_received}</code>\n"
                f"🟢 <b>Today's Total New:</b> <code>{total_day_new}</code>\n"
                f"♻️ <b>Today's Total Duplicate:</b> <code>{total_day_dup}</code>\n"
                f"🚫 <b>Today's Total Rejected:</b> <code>{total_day_rej}</code>"
            )
            
            full_text = "\n".join(blocks)
            for start_pos in range(0, len(full_text), 3500):
                await query.message.reply_text(full_text[start_pos:start_pos + 3500], parse_mode="HTML")
            return

        if data == "adm_all_ids_filter":
            date = today()
            day_data = DATA["days"].get(date)
            
            if not day_data or not day_data["ids"]:
                await query.message.reply_text(f"📅 <b>Date:</b> <code>{date}</code>\n\n📭 আজকের দিনে এখনো কোনো আইডি জমা পড়েনি।", parse_mode="HTML")
                return

            all_new_ids = []
            all_rej_ids = []
            for mem in day_data["members"].values():
                all_new_ids.extend(mem.get("new_ids", []))
                all_rej_ids.extend(mem.get("rejected_ids", []))

            text_msg = (
                "<b>📋 TODAY'S RECEIVED IDs MENU</b>\n"
                f"📅 <b>Date:</b> <code>{date}</code>\n"
                f"📦 <b>Total Received:</b> <code>{len(day_data['ids'])}</code>\n\n"
                f"🟢 <b>New IDs:</b> <code>{len(all_new_ids)}</code>\n"
                f"🚫 <b>Rejected IDs:</b> <code>{len(all_rej_ids)}</code>\n\n"
                "নিচের বাটন চেপে আপনার প্রয়োজনীয় আইডিগুলোর এক্সেল ফাইল ডাউনলোড করে নিন:"
            )

            keyboard = [
                [InlineKeyboardButton("🟢 Download Only New IDs Excel", callback_data="copy_ids_new")],
                [InlineKeyboardButton("🚫 Download Only Rejected IDs Excel", callback_data="copy_ids_rej")],
                [InlineKeyboardButton("📦 Download All Received IDs Excel", callback_data="copy_ids_all")]
            ]
            await query.message.reply_text(text_msg, reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
            return

        if data in ("copy_ids_new", "copy_ids_rej", "copy_ids_all"):
            date = today()
            day_data = DATA["days"].get(date)
            if not day_data:
                await query.message.reply_text("📭 ডাটা পাওয়া যায়নি।")
                return

            target_ids = []
            label = ""
            if data == "copy_ids_new":
                label = "NEW_IDs"
                for mem in day_data["members"].values():
                    target_ids.extend(mem.get("new_ids", []))
            elif data == "copy_ids_rej":
                label = "REJECTED_IDs"
                for mem in day_data["members"].values():
                    target_ids.extend(mem.get("rejected_ids", []))
            else:
                label = "ALL_RECEIVED_IDs"
                target_ids = list(day_data["ids"])

            if not target_ids:
                await query.message.reply_text(f"📭 আজকের দিনে কোনো আইডির রেকর্ড পাওয়া যায়নি।")
                return

            def gen_filter_excel():
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = label
                ws.append(["UID", "Date"])
                for u in target_ids:
                    ws.append([u, date])
                bio = io.BytesIO()
                wb.save(bio)
                bio.seek(0)
                return bio

            try:
                excel_bio = await asyncio.to_thread(gen_filter_excel)
                await query.message.reply_document(
                    document=excel_bio,
                    filename=f"{label.lower()}_{date}.xlsx",
                    caption=f"📋 আজকের {label.replace('_', ' ')} এর এক্সেল ফাইল (মোট: {len(target_ids)} টি)।"
                )
            except Exception as e:
                await query.message.reply_text(f"🚫 ফাইল তৈরি করা যায়নি: {e}")
            return

        if data == "adm_set_p_only":
            context.user_data["waiting_for"] = "set_price_only"
            await query.message.reply_text("✏️ <b>ONLY UID</b> এর প্রতি নতুন আইডির প্রাইস (BDT) লিখে পাঠান (যেমন: 1.5):", parse_mode="HTML")
            return

        if data == "adm_set_p_page":
            context.user_data["waiting_for"] = "set_price_page"
            await query.message.reply_text("✏️ <b>PAGE UID</b> এর প্রতি নতুন আইডির প্রাইস (BDT) লিখে পাঠান (যেমন: 2.0):", parse_mode="HTML")
            return

        if data == "adm_edit_pay_msg":
            context.user_data["waiting_for"] = "set_pay_template"
            cur_tpl = DATA["settings"].get("payment_msg_template", "")
            await query.message.reply_text(
                f"✉️ <b>EDIT PAYMENT MESSAGE TEMPLATE</b>\n\n"
                f"বর্তমান টেমপ্লেট:\n<code>{cur_tpl}</code>\n\n"
                f"আপনার নতুন কাস্টম মেসেজ লিখে পাঠান। টাকা সংক্রান্ত স্থানে <code>{{amount}}</code> এবং তারিখের স্থানে <code>{{date}}</code> ট্যাগ ব্যবহার করুন।",
                parse_mode="HTML"
            )
            return

        if data == "adm_view_custom_date":
            context.user_data["waiting_for"] = "view_custom_date"
            await query.message.reply_text("📅 যে তারিখের রিপোর্ট দেখতে চান সেই তারিখটি YYYY-MM-DD ফরম্যাটে লিখুন (যেমন: <code>2026-10-18</code>):", parse_mode="HTML")
            return

        if data == "adm_export_custom_date_btn":
            context.user_data["waiting_for"] = "export_custom_date"
            await query.message.reply_text("📅 যে তারিখের এক্সেল রিপোর্ট ডাউনলোড করতে চান সেই তারিখটি YYYY-MM-DD ফরম্যাটে লিখুন (যেমন: <code>2026-10-18</code>):", parse_mode="HTML")
            return

        if data == "adm_reject_bad_ids":
            context.user_data["waiting_for"] = "reject_bad_ids"
            await query.message.reply_text(
                "🚫 <b>REJECT SPECIFIC IDs</b>\n\n"
                "যে নষ্ট বা রিজেক্ট করা আইডিগুলো মেম্বারদের ফেরত পাঠাতে চান, সেগুলো এক বা একাধিক লাইনে লিখে মেসেজে পাঠান:",
                parse_mode="HTML"
            )
            return

        if data.startswith("acc_"):
            sub_id = data.replace("acc_", "")
            sub_data = DATA["pending_submissions"].get(sub_id)

            if not sub_data:
                try:
                    await query.edit_message_text("⚠️ এই সাবমিশনটি ইতিমধ্যে প্রসেস বা বাতিল করা হয়েছে।")
                except Exception:
                    pass
                return

            class DummyUser:
                def __init__(self, uid, name):
                    self.id = uid
                    self.full_name = name

            d_user = DummyUser(sub_data["user_id"], sub_data["full_name"])
            cat_type = sub_data.get("category", "ONLY UID")

            await process_valid_submission(
                d_user,
                sub_data["ids"],
                cat_type,
                context,
                query.message,
                sub_data.get("file_info")
            )

            del DATA["pending_submissions"][sub_id]
            save_db()

            try:
                await query.edit_message_text(f"🟢 Approved & Notified!\n👤 Name: {html.escape(sub_data['full_name'])} ({len(sub_data['ids'])} IDs Batch)")
            except Exception:
                pass
            return

        if data.startswith("rej_"):
            sub_id = data.replace("rej_", "")
            sub_data = DATA["pending_submissions"].get(sub_id)

            if sub_data:
                u_id = sub_data["user_id"]
                try:
                    await context.bot.send_message(
                        chat_id=u_id,
                        text="🚫 <b>আপনার পাঠানো সাবমিশনটি অ্যাডমিন কর্তৃক বাতিল (Rejected) করা হয়েছে।</b>",
                        parse_mode="HTML"
                    )
                except Exception:
                    pass

                del DATA["pending_submissions"][sub_id]
                save_db()

            try:
                await query.edit_message_text("🚫 Submissions Rejected! (মেম্বারের কাছে নোটিফিকেশন পাঠানো হয়েছে)")
            except Exception:
                pass
            return

        if data == "adm_all_members":
            all_members = {}
            for d_val in DATA["days"].values():
                for uid, mem in d_val["members"].items():
                    if uid not in all_members:
                        all_members[uid] = mem

            if not all_members:
                await query.message.reply_text("📭 ডাটাবেজে এখনো কোনো মেম্বারের রেকর্ড নেই।")
                return

            lines = [f"👥 <b>বট ব্যবহারকারী সকল মেম্বারদের তালিকা (মোট: {len(all_members)} জন):</b>\n"]
            for uid, mem in all_members.items():
                lines.append(f"👤 Name: <a href='tg://user?id={uid}'>{html.escape(mem['name'])}</a>\n🆔 User ID: <code>{uid}</code>\n------------------------------")

            full_output = "\n".join(lines)
            for start_pos in range(0, len(full_output), 3500):
                await query.message.reply_text(full_output[start_pos:start_pos + 3500], parse_mode="HTML")
            return

        if data == "adm_reset_all":
            DATA["days"] = {}
            DATA["pending_submissions"] = {}
            DATA["payments"] = {}
            save_db()
            await query.message.reply_text("🔥 <b>সফলভাবে সমস্ত আইডি, পেমেন্ট ও সাবমিশন হিস্ট্রি রিসেট করা হয়েছে!</b>", parse_mode="HTML")
            return

        if data == "adm_user_hist":
            context.user_data["waiting_for"] = "view_user_hist"
            await query.message.reply_text("👤 যে ইউজারের হিস্ট্রি দেখতে চান তার <b>User ID</b> লিখে পাঠান:", parse_mode="HTML")
            return

        if data == "adm_del_hist":
            context.user_data["waiting_for"] = "del_hist_type"
            await query.message.reply_text(
                "🗑️ <b>DELETE HISTORY OPTIONS</b>\n\n"
                "• কোনো নির্দিষ্ট তারিখের হিস্ট্রি মুছতে চাইলে তারিখ লিখুন (যেমন: <code>2026-10-13</code>)\n"
                "• নির্দিষ্ট কোনো ইউজারের অল-টাইম ডাটা মুছতে তার <b>User ID</b> লিখুন:",
                parse_mode="HTML"
            )
            return

        if data == "adm_msg_user":
            context.user_data["waiting_for"] = "msg_user_id"
            await query.message.reply_text("✉️ <b>MESSAGE SPECIFIC USER</b>\n\nযে ইউজারের কাছে মেসেজ পাঠাতে চান তার <b>User ID</b> লিখে পাঠান:", parse_mode="HTML")
            return

        if data == "adm_view":
            s = DATA["settings"]
            text = (
                "📋 <b>SYSTEM CONFIGURATION</b>\n\n"
                f"📢 Channels: <code>{s.get('channels')}</code>\n"
                f"👑 Admin IDs: <code>{s.get('admin_ids')}</code>\n"
                f"🎯 Receivers: <code>{s['target_chat_ids']}</code>\n"
                f"🔢 Min IDs Limit: <code>{s['min_ids_limit']}</code>\n"
                f"🛡️ Admin Approval Mode: <code>{s['approval_mode']}</code>\n"
                f"🛠️ Maintenance Mode: <code>{s['maintenance']}</code>\n"
                f"⚙️ Duplicate Check: <code>{s['duplicate_mode']}</code>\n"
                f"⏰ Work Schedule: <code>{s.get('work_schedule_enabled')}</code> ({s.get('start_time')} - {s.get('end_time')})\n"
                f"📁 Live Check File: <code>{s.get('send_live_check_file', True)}</code>\n"
                f"💵 Price ONLY UID: <code>{s.get('price_only_uid', 0.0)} BDT</code>\n"
                f"📄 Price PAGE UID: <code>{s.get('price_page_uid', 0.0)} BDT</code>\n"
                f"🚫 Banned Users: <code>{len(DATA['banned_users'])}</code>"
            )
            await query.message.reply_text(text, parse_mode="HTML")
            return

        if data == "adm_toggle_approval":
            DATA["settings"]["approval_mode"] = not DATA["settings"]["approval_mode"]
            save_db()
            await render_sys_settings(query)
            return

        if data == "adm_toggle_maint":
            DATA["settings"]["maintenance"] = not DATA["settings"]["maintenance"]
            save_db()
            await render_sys_settings(query)
            return

        if data == "adm_toggle_mode":
            DATA["settings"]["duplicate_mode"] = "all_time" if DATA["settings"]["duplicate_mode"] == "daily" else "daily"
            save_db()
            await render_sys_settings(query)
            return

        if data == "adm_add_rcv":
            context.user_data["waiting_for"] = "add_rcv"
            await query.message.reply_text("📥 নতুন রিসিভার চ্যাট আইডি (Chat ID) লিখে পাঠান:")
            return

        if data == "adm_rem_rcv":
            context.user_data["waiting_for"] = "rem_rcv"
            await query.message.reply_text(f"📤 বাদ দেওয়ার জন্য আইডি পাঠান। বর্তমান লিস্ট: <code>{DATA['settings']['target_chat_ids']}</code>", parse_mode="HTML")
            return

        if data == "adm_ban":
            context.user_data["waiting_for"] = "ban_user"
            await query.message.reply_text("🚫 ব্যান বা আনব্যান করার জন্য ইউজারের User ID পাঠান:")
            return

        if data == "adm_export":
            def generate_excel():
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Submissions Report"
                ws.append(["Date", "User ID", "Name", "Received", "New", "Only UID", "Page UID", "Duplicate", "Rejected", "Earned BDT", "Payment Info", "Payment Status"])

                for d_date, d_val in DATA["days"].items():
                    for u_id, mem in d_val["members"].items():
                        if mem["received"] > 0:
                            earned = calculate_member_earned(mem)
                            pay_st = get_user_payment_status(u_id, d_date)
                            pay_dets = DATA["user_payment_details"].get(str(u_id), {})
                            pay_str = ", ".join([f"{m}: {n}" for m, n in pay_dets.items()]) if isinstance(pay_dets, dict) and pay_dets else "Not Set"
                            ws.append([d_date, u_id, mem["name"], mem["received"], mem["new"], mem.get("only_uid_count",0), mem.get("page_uid_count",0), mem["duplicate"], mem.get("rejected", 0), earned, pay_str, pay_st])

                bio = io.BytesIO()
                wb.save(bio)
                bio.seek(0)
                return bio

            excel_bio = await asyncio.to_thread(generate_excel)
            await query.message.reply_document(
                document=excel_bio,
                filename=f"report_all_{today()}.xlsx",
                caption="📊 সম্পূর্ণ সাবমিশনের এক্সেল রিপোর্ট (পেমেন্ট নম্বরসহ)।"
            )
            return
    except Exception as e:
        print(f"Admin Callback Error: {e}")


# ================= MAIN APP & HANDLERS =================

def main():
    try:
        loop = asyncio.get_event_loop()
        if loop.is_closed():
            raise RuntimeError("Loop is closed")
    except RuntimeError:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
    
    app = (
        Application.builder()
        .token(BOT_TOKEN)
        .read_timeout(30)
        .write_timeout(30)
        .connect_timeout(30)
        .pool_timeout(30)
        .build()
    )

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("status", status))
    app.add_handler(CommandHandler("myids", my_ids_command))
    app.add_handler(CommandHandler("setpayment", set_payment))
    app.add_handler(CommandHandler("language", language_command))
    app.add_handler(CommandHandler("support", support_command))
    app.add_handler(CommandHandler("admin", admin))
    app.add_handler(CallbackQueryHandler(admin_callback))

    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, receive_text))
    app.add_handler(MessageHandler(filters.Document.ALL, receive_text))

    print("🤖 Ultra-High-Tech Single Telegram Bot is running smoothly...")
    app.run_polling(drop_pending_updates=True)


from keep_alive import keep_alive

if __name__ == "__main__":
    keep_alive()
    main()
